"""Unit tests for utils.extract — text out of attachments. No Outlook needed.

The Office test files are generated with the same libraries that read them, so these
tests also prove the libraries are installed.
"""
import pathlib
import tempfile
import unittest

from outlook_desktop_mcp.utils import extract as x


class PageSpecTests(unittest.TestCase):
    def test_parse_pages(self):
        self.assertEqual(x.parse_pages("", 4), [0, 1, 2, 3])
        self.assertEqual(x.parse_pages("1-2,4", 4), [0, 1, 3])
        self.assertEqual(x.parse_pages("3-", 4), [2, 3])
        self.assertEqual(x.parse_pages("9", 4), [])


class ExtractTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = pathlib.Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_docx_paragraphs_headings_and_tables(self):
        import docx
        d = docx.Document()
        d.add_heading("Lease Summary", level=1)
        d.add_paragraph("Assignment requires landlord consent.")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Term"
        t.cell(0, 1).text = "5 years"
        t.cell(1, 0).text = "Rent"
        t.cell(1, 1).text = "$4,000"
        p = self.dir / "lease.docx"
        d.save(str(p))
        text, info = x.extract(p)
        self.assertEqual(info["kind"], "docx")
        self.assertIn("# Lease Summary", text)
        self.assertIn("Assignment requires landlord consent.", text)
        self.assertIn("Term | 5 years", text)
        self.assertEqual(info["tables"], 1)
        self.assertFalse(info["truncated"])

    def test_xlsx_sheets_rows_and_sheet_filter(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "QoE"
        ws.append(["Line", "FY24", "FY25"])
        ws.append(["Revenue", 1000, 1200])
        ws2 = wb.create_sheet("Adjustments")
        ws2.append(["Owner comp", 150])
        p = self.dir / "qoe.xlsx"
        wb.save(str(p))
        text, info = x.extract(p)
        self.assertEqual(info["sheets"], ["QoE", "Adjustments"])
        self.assertIn("--- sheet: QoE ---", text)
        self.assertIn("Revenue\t1000\t1200", text)
        self.assertIn("Owner comp\t150", text)
        text2, info2 = x.extract(p, sheet="adjust")
        self.assertEqual(info2["sheets_read"], ["Adjustments"])
        self.assertNotIn("Revenue", text2)
        _, info3 = x.extract(p, sheet="nope")
        self.assertIn("error", info3)

    def test_pptx_slides_and_page_selection(self):
        from pptx import Presentation
        prs = Presentation()
        for title in ("Overview", "Financials", "Team"):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = title
            slide.placeholders[1].text = f"Body of {title}"
        p = self.dir / "cim.pptx"
        prs.save(str(p))
        text, info = x.extract(p)
        self.assertEqual(info["slides_total"], 3)
        self.assertIn("--- slide 2 ---\nFinancials\nBody of Financials", text)
        text2, info2 = x.extract(p, pages="3")
        self.assertEqual(info2["slides_read"], [3])
        self.assertNotIn("Overview", text2)

    def test_pdf_without_text_layer_is_reported(self):
        from pypdf import PdfWriter
        w = PdfWriter()
        w.add_blank_page(width=200, height=200)
        p = self.dir / "scan.pdf"
        with open(p, "wb") as f:
            w.write(f)
        text, info = x.extract(p)
        self.assertEqual(info["pages_total"], 1)
        self.assertIn("scanned", info.get("note", ""))

    def test_csv_text_and_cap(self):
        p = self.dir / "list.csv"
        p.write_text("name,phone\nChad,555\nLinda,556\n", encoding="utf-8")
        text, info = x.extract(p)
        self.assertEqual(info["rows_read"], 3)
        self.assertIn("Chad\t555", text)
        big = self.dir / "big.txt"
        big.write_text("x" * 5000, encoding="utf-8")
        text, info = x.extract(big, max_chars=1000)
        self.assertEqual(len(text), 1000)
        self.assertTrue(info["truncated"])

    def test_unsupported_extension(self):
        p = self.dir / "photo.png"
        p.write_bytes(b"\x89PNG")
        text, info = x.extract(p)
        self.assertEqual(text, "")
        self.assertIn("Cannot extract", info["error"])


if __name__ == "__main__":
    unittest.main()
